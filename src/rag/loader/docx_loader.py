"""
Word 文档加载器
"""

from pathlib import Path
from docx import Document as DocxDocument
from .base import BaseLoader, Document


class DocxLoader(BaseLoader):
    def load(self, source: str | Path) -> list[Document]:
        path = Path(source)
        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n".join(paragraphs)
        return [Document(
            content=content,
            metadata={
                "source": str(path),
                "type": "document",
                "filename": path.name,
            }
        )]
